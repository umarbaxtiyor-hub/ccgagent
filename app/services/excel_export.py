from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models import Transaction, TransactionType


async def build_report(session: AsyncSession, start: date, end: date) -> BytesIO:
    result = await session.execute(
        select(Transaction)
        .where(Transaction.occurred_on >= start, Transaction.occurred_on <= end)
        .options(
            selectinload(Transaction.category),
            selectinload(Transaction.created_by),
            selectinload(Transaction.project),
        )
        .order_by(Transaction.occurred_on, Transaction.id)
    )
    transactions = result.scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Tranzaksiyalar"
    headers = ["Sana", "Loyiha", "Turi", "Kategoriya", "Summasi", "Tavsif", "Kontragent", "Kim kiritdi"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    total_income = 0.0
    total_expense = 0.0
    for t in transactions:
        type_label = "Kirim" if t.type == TransactionType.income else "Chiqim"
        amount = float(t.amount)
        if t.type == TransactionType.income:
            total_income += amount
        else:
            total_expense += amount
        ws.append(
            [
                t.occurred_on.isoformat(),
                t.project.name if t.project else "-",
                type_label,
                t.category.name if t.category else "",
                amount,
                t.description,
                t.counterparty,
                t.created_by.full_name if t.created_by else "",
            ]
        )

    ws.append([])
    ws.append(["", "", "", "Jami kirim", total_income])
    ws.append(["", "", "", "Jami chiqim", total_expense])
    ws.append(["", "", "", "Balans", total_income - total_expense])

    summary_ws = wb.create_sheet("Kategoriyalar bo'yicha")
    summary_ws.append(["Loyiha", "Turi", "Kategoriya", "Jami summa"])
    for cell in summary_ws[1]:
        cell.font = Font(bold=True)

    totals_by_category: dict[tuple[str, str, str], float] = {}
    for t in transactions:
        type_label = "Kirim" if t.type == TransactionType.income else "Chiqim"
        project_name = t.project.name if t.project else "-"
        key = (project_name, type_label, t.category.name if t.category else "Noma'lum")
        totals_by_category[key] = totals_by_category.get(key, 0.0) + float(t.amount)

    for (project_name, type_label, category_name), total in sorted(totals_by_category.items()):
        summary_ws.append([project_name, type_label, category_name, total])

    for sheet in (ws, summary_ws):
        for column_cells in sheet.columns:
            length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 40)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
